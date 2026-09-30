"""Build a styled workbook containing one account's saved job-board data."""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ILLEGAL_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
STATUS_COLORS = {
    '关注': ('FFF5DF', '855A12'),
    '已投递': ('EAF3FC', '286997'),
    '笔试': ('E6F6F4', '24777B'),
    '面试': ('F2ECFA', '70509B'),
    'Offer': ('E5F4EA', '256D47'),
    '流程终止': ('FDECEB', 'A83737'),
    '暂不考虑': ('F0F1F0', '65726A'),
}


def cell_text(value):
    """Keep arbitrary source text safe as a literal Excel string."""
    return ILLEGAL_XML.sub('', str(value if value is not None else ''))[:32767]


def add_sheet(book, name, columns, rows, widths, status_column=None):
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'C2' if len(columns) > 2 else 'A2'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.print_title_rows = '1:1'
    sheet.append(columns)
    sheet.row_dimensions[1].height = 30
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='176456')
        cell.font = Font(name='Microsoft YaHei', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center')
    for row_index, values in enumerate(rows, 2):
        sheet.row_dimensions[row_index].height = 30
        for column_index, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column_index, cell_text(value))
            cell.data_type = 's'
            cell.font = Font(name='Microsoft YaHei', size=10, color='253C32')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if row_index % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F5F8F5')
            if column_index == status_column and value in STATUS_COLORS:
                background, foreground = STATUS_COLORS[value]
                cell.fill = PatternFill('solid', fgColor=background)
                cell.font = Font(name='Microsoft YaHei', size=10, color=foreground, bold=True)
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}'


def build_workbook(annotations, bookmarks):
    book = Workbook()
    book.remove(book.active)
    book.properties.title = '个人求职数据'
    book.properties.creator = '大潘的就业情报站'
    add_sheet(book, '投递跟进',
              ['公司', '岗位', '投递流程', '我的标签', '我的备注', '工作地点', '网申截止',
               '网申公告', '投递链接', '最后修改', '资料表ID', '记录ID'],
              annotations, [25, 38, 16, 24, 54, 23, 18, 48, 48, 22, 25, 25], 3)
    add_sheet(book, '我的收藏',
              ['名称', '类型', '岗位 / 内容', '地点', '时间 / 截止', '来源链接', '备注', '收藏时间', '记录ID'],
              bookmarks, [30, 14, 42, 25, 25, 48, 48, 22, 28])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
